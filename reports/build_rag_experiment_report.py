from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "667085"
LIGHT_GRAY = "F2F4F7"
BLUE_GRAY = "E8EEF5"
LIGHT_GREEN = "E2F0D9"
LIGHT_GOLD = "FFF2CC"
BORDER = "CBD5E1"


def set_run_font(
    run,
    *,
    name: str = "Calibri",
    east_asia: str = "맑은 고딕",
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, size: float, color: str = "000000", bold: bool = False) -> None:
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "맑은 고딕")


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    set_style_font(normal, 11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        set_style_font(style, size, color, True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = doc.styles["Title"]
    set_style_font(title, 24, INK, True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(6)

    subtitle = doc.styles["Subtitle"]
    set_style_font(subtitle, 12.5, MUTED, False)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(14)

    if "Code Block" not in doc.styles:
        code = doc.styles.add_style("Code Block", 1)
    else:
        code = doc.styles["Code Block"]
    set_style_font(code, 9, "263238", False)
    code.font.name = "Consolas"
    rpr = code.element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), "Consolas")
    rpr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.left_indent = Inches(0.12)
    code.paragraph_format.right_indent = Inches(0.12)


def set_page_geometry(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)


def add_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))


def configure_header_footer(doc: Document) -> None:
    for section in doc.sections:
        header = section.header
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run("SAFEMAINT AI  |  RAG SAMPLE EXPERIMENT")
        set_run_font(run, size=8.5, color=MUTED, bold=True)

        footer = section.footer
        footer_paragraph = footer.paragraphs[0]
        footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer_paragraph.paragraph_format.space_before = Pt(0)
        footer_paragraph.paragraph_format.space_after = Pt(0)
        label = footer_paragraph.add_run("SafeMaint AI  •  ")
        set_run_font(label, size=8.5, color=MUTED)
        page_run = footer_paragraph.add_run()
        set_run_font(page_run, size=8.5, color=MUTED)
        add_field(page_run, "PAGE")
        middle = footer_paragraph.add_run(" / ")
        set_run_font(middle, size=8.5, color=MUTED)
        total_run = footer_paragraph.add_run()
        set_run_font(total_run, size=8.5, color=MUTED)
        add_field(total_run, "NUMPAGES")


def shade(element, fill: str) -> None:
    shd = element.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        element.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in CELL_MARGINS_DXA.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), BORDER)


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must total {CONTENT_WIDTH_DXA}: {widths_dxa}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            cell.width = Twips(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_table_borders(table)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    color: str = "000000",
    align=WD_ALIGN_PARAGRAPH.LEFT,
    size: float = 9.5,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.10
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths_dxa: list[int],
    *,
    status_column: int | None = None,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    repeat_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell._tc.get_or_add_tcPr(), LIGHT_GRAY)
        set_cell_text(cell, header, bold=True, color=INK, size=9.5)

    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            set_cell_margins(cells[index])
            cells[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            align = (
                WD_ALIGN_PARAGRAPH.CENTER
                if index == status_column
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            set_cell_text(cells[index], value, align=align)
            if status_column is not None and index == status_column:
                fill = LIGHT_GREEN if value in {"통과", "정상", "완료"} else LIGHT_GOLD
                shade(cells[index]._tc.get_or_add_tcPr(), fill)
    set_table_geometry(table, widths_dxa)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(2)


def create_numbering(doc: Document, num_format: str, level_text: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), num_format)
    level.append(fmt)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), level_text)
    level.append(text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    level.append(suffix)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    ppr.append(ind)
    level.append(ppr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_item(doc: Document, text: str, num_id: int) -> None:
    paragraph = doc.add_paragraph()
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend((level, num))
    ppr.append(num_pr)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.167
    run = paragraph.add_run(text)
    set_run_font(run, size=11)


def add_callout(doc: Document, label: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.10
    ppr = paragraph._p.get_or_add_pPr()
    shade(ppr, BLUE_GRAY)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), BLUE)
    borders.append(left)
    ppr.append(borders)
    label_run = paragraph.add_run(f"{label}  ")
    set_run_font(label_run, size=11, color=INK, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=11, color=INK)


def add_code_block(doc: Document, code: str) -> None:
    paragraph = doc.add_paragraph(style="Code Block")
    ppr = paragraph._p.get_or_add_pPr()
    shade(ppr, LIGHT_GRAY)
    run = paragraph.add_run(code)
    set_run_font(run, name="Consolas", east_asia="맑은 고딕", size=9)


def add_heading(doc: Document, text: str, level: int) -> None:
    doc.add_paragraph(text, style=f"Heading {level}")


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, size=11, bold=True, color=INK)
        rest = paragraph.add_run(text[len(bold_prefix) :])
        set_run_font(rest, size=11)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=11)


def build_report(output_path: Path) -> None:
    doc = Document()
    configure_styles(doc)
    set_page_geometry(doc)
    configure_header_footer(doc)
    doc.core_properties.title = "SafeMaint 사고 데이터 적재 및 BGE-M3 임베딩 실험 결과"
    doc.core_properties.subject = "격리 DB 기반 RAG 샘플 실험"
    doc.core_properties.author = "SafeMaint AI Team"
    doc.core_properties.keywords = "SafeMaint, BGE-M3, pgvector, RAG, PostgreSQL"

    bullet_id = create_numbering(doc, "bullet", "•")
    decimal_id = create_numbering(doc, "decimal", "%1.")

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(8)
    kicker.paragraph_format.space_after = Pt(4)
    run = kicker.add_run("TECHNICAL EXPERIMENT REPORT")
    set_run_font(run, size=9.5, color=BLUE, bold=True)

    doc.add_paragraph(
        "SafeMaint 사고 데이터 적재 및\nBGE-M3 임베딩 실험 결과",
        style="Title",
    )
    doc.add_paragraph(
        "PostgreSQL · pgvector · 격리 테스트 DB 검증",
        style="Subtitle",
    )
    metadata = (
        ("프로젝트", "SafeMaint AI"),
        ("대상 브랜치", "feature/chan"),
        ("기준 코드", "dev 66541c0 이후 신규 구현"),
        ("실험 일자", "2026-07-16 (Asia/Seoul)"),
        ("격리 DB", "safemaint_rag_test"),
        ("모델", "BAAI/bge-m3"),
    )
    for label, value in metadata:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        label_run = paragraph.add_run(f"{label}: ")
        set_run_font(label_run, size=10.5, color=INK, bold=True)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, size=10.5)

    add_callout(
        doc,
        "결론",
        "요청된 1~8단계를 모두 완료했다. 기존 safemaint 개발 DB와 Docker 볼륨을 "
        "삭제하지 않았으며, 격리 DB에 문서 200개·청크 245개를 멱등 적재하고 "
        "BGE-M3 1024차원 임베딩 245개 및 5개 질의 코사인 검색을 검증했다.",
    )

    add_heading(doc, "1. 경영진 요약", 1)
    add_table(
        doc,
        ["검증 항목", "결과", "판정"],
        [
            ["원본 데이터 dry-run", "9,223문서 / 12,296청크 전수 스캔", "통과"],
            ["샘플 적재", "domestic 100 + fatal 100 / 청크 245", "완료"],
            ["재실행 멱등성", "문서 200·청크 245 모두 unchanged", "통과"],
            ["BGE-M3 임베딩", "ready 245 / failed 0 / skipped 0", "완료"],
            ["벡터 일관성", "단일 모델·1024차원·NULL/불일치 0", "통과"],
            ["pgvector 검색", "5개 질의 실행, 임계값 0.25", "통과"],
            ["회귀 테스트", "Backend 27 / AI 7 / Frontend build", "통과"],
        ],
        [2700, 5220, 1440],
        status_column=2,
    )
    add_body(
        doc,
        "핵심 판단: 데이터 적재와 벡터 저장·조회 경로는 다음 실험으로 확장 가능한 "
        "상태다. 다만 첫 200개 문서로 구성한 표본은 일부 질의의 사건 유형을 충분히 "
        "포함하지 않아 검색 품질을 전체 데이터 품질로 일반화하면 안 된다.",
        bold_prefix="핵심 판단:",
    )

    add_heading(doc, "2. 실험 범위와 안전 통제", 1)
    for item in (
        "실제 쓰기는 이름이 _test로 끝나는 safemaint_rag_test에만 허용했다. CLI와 서비스 계층에서 이 조건을 이중 검사한다.",
        "기존 safemaint 개발 DB, Docker named volume, .env, 전처리 원본 JSONL은 수정하거나 삭제하지 않았다.",
        "docker compose down -v, docker volume rm, truncate, 전체 DB 삭제 명령은 사용하지 않았다.",
        "전체 9,223개 문서와 12,296개 청크 적재, BM25, RRF, Reranker, LLM 답변 생성은 범위에서 제외했다.",
        "데이터 이용조건 확정 전 모든 신규 documents.access_level은 restricted로 저장했다.",
    ):
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "3. 원본 데이터 점검 및 dry-run", 1)
    add_table(
        doc,
        ["점검 항목", "관측값", "해석"],
        [
            ["accident_documents.jsonl", "9,223", "domestic 6,348 / fatal 2,875"],
            ["accident_chunks.jsonl", "12,296", "전체 스트리밍 검사"],
            ["선택 문서", "200", "source별 최초 100개"],
            ["선택 청크", "245", "선택 문서에 연결된 모든 청크"],
            ["빈 제목 / 빈 청크", "0 / 0", "적재 중단 사유 없음"],
            ["중복 external_id / chunk index", "0 / 0", "고유성 문제 없음"],
            ["연결 불가 document_id", "0", "참조 무결성 문제 없음"],
            ["title_only 문서", "72", "본문 없이 제목·검색 텍스트 중심"],
            ["원본 메타데이터 누락", "0", "domestic/fatal 원본과 모두 연결"],
        ],
        [3000, 1800, 4560],
    )
    add_body(
        doc,
        "발견된 데이터 규격 주의사항: 국내 데이터에는 실제 null과 문자열 \"null\"이 "
        "혼재한다. 적재 계층은 둘을 동일한 빈 값으로 정규화하고, content_raw 전체를 "
        "metadata에 중복 저장하지 않는다.",
        bold_prefix="발견된 데이터 규격 주의사항:",
    )

    add_heading(doc, "4. DB 적재 매핑", 1)
    add_table(
        doc,
        ["대상", "필드", "저장 규칙"],
        [
            ["documents", "external_id", "incident:{source_type}:{source_row_id}"],
            ["documents", "source_type / access_level", "incident / restricted"],
            ["documents", "publisher·URL·revision", "근거가 없으므로 NULL"],
            ["documents", "published_at", "날짜 문자열이 명확한 경우만 Date 저장"],
            ["documents", "metadata", "dataset, 원본 ID, 업종/장소, content_quality"],
            ["document_chunks", "document_id / chunk_index", "DB UUID 연결 / chunk_order"],
            ["document_chunks", "content_hash", "CRLF→LF, 양끝 공백 제거 후 UTF-8 SHA-256"],
            ["document_chunks", "embedding 초기값", "NULL / model NULL / dimension NULL / pending"],
        ],
        [1800, 2700, 4860],
    )
    add_body(
        doc,
        "멱등 upsert는 documents.external_id와 (document_id, chunk_index)를 사용한다. "
        "동일 입력은 업데이트하지 않으며, 청크 내용 해시가 바뀌면 기존 임베딩을 "
        "NULL로 되돌리고 pending 상태로 재설정한다.",
    )

    add_heading(doc, "5. 샘플 적재와 멱등성 결과", 1)
    add_table(
        doc,
        ["실행", "문서", "청크", "결과"],
        [
            ["최초 적재", "insert 200", "insert 245", "1개 문서 batch + 1개 청크 batch 커밋"],
            ["동일 입력 2회차", "unchanged 200", "unchanged 245", "추가·갱신 0"],
            ["임베딩 후 재적재", "unchanged 200", "unchanged 245", "ready 벡터 보존"],
        ],
        [1800, 1800, 1800, 3960],
    )
    add_callout(
        doc,
        "멱등성 판정",
        "동일 데이터 재실행 시 문서·청크 수가 증가하지 않았고 unchanged로 분류됐다. "
        "동일 내용의 ready 임베딩도 삭제하거나 재처리하지 않았다.",
    )

    add_heading(doc, "6. BGE-M3 임베딩 실험", 1)
    add_table(
        doc,
        ["항목", "실제 값"],
        [
            ["모델", "BAAI/bge-m3"],
            ["실행 장치", "NVIDIA GeForce RTX 4070 Ti (CUDA)"],
            ["Torch", "2.13.0+cu130"],
            ["Sentence Transformers", "5.6.0"],
            ["처리 청크", "245 (요청 상한 500)"],
            ["상태", "ready 245 / failed 0 / skipped 0 / pending 0"],
            ["실제 벡터 차원", "1024"],
            ["정규화", "문서·질의 모두 L2 정규화, 비정상 norm 차단"],
            ["캐시", "ai/.model-cache, 4.25GB / 23 files, Git 제외"],
        ],
        [2700, 6660],
    )
    add_body(
        doc,
        "DB 검증에서 ready이면서 embedding이 NULL인 행, 모델 누락, 저장 차원과 "
        "vector_dims() 불일치, embedding_error 기록은 모두 0개였다. 동일 임베딩 명령 "
        "재실행은 processed 0으로 종료됐다.",
    )

    add_heading(doc, "7. pgvector 검색 결과", 1)
    add_body(
        doc,
        "검색은 BAAI/bge-m3·1024차원·ready 행만 대상으로 cosine distance를 계산했다. "
        "유사도 0.25 미만은 제외하고 최대 5개만 반환했다.",
    )
    add_table(
        doc,
        ["질의", "대표 결과", "평가"],
        [
            [
                "컨베이어 벨트 끼임 사고",
                "incident:domestic:31 · 운전중 콘베이어 벨트와 드럼 사이 협착 · 0.645",
                "관련 사고 1위",
            ],
            [
                "크레인 코일 상차 낙하",
                "incident:domestic:4 · 천정크레인 화물 낙하 · 0.683; 목표 fatal:1은 5위 · 0.649",
                "주제 일치",
            ],
            [
                "작동유 드럼 파열",
                "incident:domestic:1 · 작동유 드럼 파열 · 0.749",
                "정확 사례 1위",
            ],
            [
                "설비 청소 중 감전",
                "1위 fatal:66 · 정화조 추락 · 0.594; 감전 키워드 청크 0개",
                "표본 부족",
            ],
            [
                "회전체 점검 중 말림",
                "fatal:16 · 가동 스크류 말려들어감은 3위 · 0.571",
                "관련 결과 있으나 1위 아님",
            ],
        ],
        [2400, 5220, 1740],
    )
    add_body(
        doc,
        "검색 해석: 데이터에 정확한 사건이 포함된 앞의 세 질의는 상위 결과가 양호했다. "
        "감전 질의는 제한 표본에 해당 키워드가 없으므로 검색기의 실패로 단정할 수 없다. "
        "회전체 질의는 관련 청크 7개가 있으나 용어 차이와 표본 규모 때문에 순위가 낮았다.",
        bold_prefix="검색 해석:",
    )

    add_heading(doc, "8. 구현 파일", 1)
    for item in (
        "backend/app/services/document_ingestion.py - 스트리밍 검사, 매핑, 조건부 upsert, 임베딩 무효화",
        "backend/app/commands/import_accidents.py - dry-run 및 샘플 적재 CLI",
        "backend/app/services/document_embeddings.py - BGE-M3 adapter, 배치 임베딩, 검색",
        "backend/app/commands/embed_document_chunks.py - 최대 500개 임베딩 CLI",
        "backend/app/commands/search_document_chunks.py - 5개 질의 cosine 검색 CLI",
        "backend/tests/test_document_ingestion.py / test_document_embeddings.py - 단위·격리 DB 테스트",
        "scripts/run-rag-experiment.ps1 - 1~8단계 재실행 자동화",
        "ai/requirements-embedding.txt - 선택 ML 의존성 및 CUDA 설치 안내",
        "README.md / ai/README.md / .gitignore - 실행법·안전 제한·캐시 제외",
    ):
        add_list_item(doc, item, bullet_id)
    add_body(
        doc,
        "기존 Alembic migration, DB 모델, 전처리 JSONL은 수정하지 않았다. 현재 Vector()는 "
        "차원 미지정 상태를 유지하고 모델명과 실제 차원 필터로 혼합 벡터를 차단했다.",
    )

    add_heading(doc, "9. 테스트와 검증", 1)
    add_table(
        doc,
        ["검증", "결과", "비고"],
        [
            ["Backend pytest", "27 passed", "safemaint_rag_test 통합 테스트 포함"],
            ["AI pytest", "7 passed", "PDF 전처리 회귀"],
            ["Embedding pip check", "No broken requirements", "선택 환경 정상"],
            ["Frontend build", "통과", "Node 24.14 / Next 16.2.10 / TypeScript"],
            ["Docker Compose config", "통과", "기존 서비스 설정 유지"],
            ["PowerShell parse", "통과", "Windows PowerShell 5 UTF-8 BOM 호환"],
            ["Python compileall / git diff --check", "통과", "구문·공백 오류 없음"],
            ["DB SQL consistency", "통과", "NULL·모델·차원·상태 검사"],
        ],
        [2700, 2340, 4320],
    )
    add_body(
        doc,
        "비차단 경고: 기존 FastAPI 테스트 환경에서 Starlette의 httpx TestClient 사용에 대한 "
        "deprecation warning이 1건 발생했다. 이번 변경으로 생긴 실패는 아니다.",
        bold_prefix="비차단 경고:",
    )

    add_heading(doc, "10. 한계와 전체 적재 전 해결사항", 1)
    limitations = (
        "표본 편향: 최초 100개씩만 사용해 감전 사례가 0개였다. 질의별 층화 표본 또는 더 넓은 검증셋이 필요하다.",
        "본문 품질: 200문서 중 72개가 title_only다. 검색 평가는 full/title_only를 분리해 측정해야 한다.",
        "벡터 인덱스: 245개는 순차 검색으로 충분하지만 전체 적재 전 모델을 확정하고 새 Alembic revision으로 고정 차원 및 HNSW를 검토해야 한다.",
        "임계값: 0.25는 동작 확인값이다. 정답셋을 만든 뒤 precision/recall 기반으로 재조정해야 한다.",
        "Windows 캐시: 심볼릭 링크가 비활성화돼 모델 캐시가 4.25GB다. 필요 시 개발자 모드를 검토하되 필수는 아니다.",
        "Node 환경: 시스템 Node 22.11은 frontend 요구조건(>=22.13)을 충족하지 않는다. 팀 환경은 Node를 업데이트해야 한다.",
        "라이선스·데이터 정책: 사고 데이터 이용조건과 공용 staging DB 정책을 확정하기 전 restricted를 유지한다.",
        "다음 범위: BM25, RRF, Reranker, LLM 답변 생성과 전체 임베딩은 별도 승인 후 진행한다.",
    )
    for item in limitations:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "11. 재실행 방법", 1)
    add_body(doc, "프로젝트 루트에서 다음 순서로 실행한다. 실제 비밀번호는 명령이나 Git에 기록하지 않는다.")
    for step in (
        "feature/chan 최신 코드와 .env를 준비한다.",
        "임베딩 전용 가상환경과 선택 의존성을 설치한다.",
        "NVIDIA GPU 사용 시 공식 PyTorch CUDA wheel로 Torch를 교체한다.",
        "자동화 스크립트를 실행한다. 스크립트는 _test DB만 허용하고 기존 DB를 삭제하지 않는다.",
    ):
        add_list_item(doc, step, decimal_id)
    add_code_block(
        doc,
        "git switch feature/chan\n"
        "git pull origin feature/chan\n\n"
        "python -m venv ai\\.venv-embedding\n"
        ".\\ai\\.venv-embedding\\Scripts\\python.exe -m pip install `\n"
        "  -r ai\\requirements-embedding.txt\n\n"
        ".\\ai\\.venv-embedding\\Scripts\\python.exe -m pip install `\n"
        "  \"torch==2.13.0+cu130\" `\n"
        "  --index-url https://download.pytorch.org/whl/cu130 --force-reinstall\n\n"
        ".\\scripts\\run-rag-experiment.ps1",
    )

    add_heading(doc, "12. 최종 판정", 1)
    add_callout(
        doc,
        "판정: 성공",
        "격리 DB에서 적재·멱등성·임베딩·검색·회귀 테스트가 모두 완료됐다. "
        "현재 결과는 기술 경로의 작동 증명이며 전체 데이터 검색 품질의 최종 평가는 아니다. "
        "전체 적재와 인덱스 설계는 위 한계를 해결한 뒤 다음 승인 단계에서 진행한다.",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build_report(args.output.resolve())
    print(args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
